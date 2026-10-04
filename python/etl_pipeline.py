import pandas as pd
import mysql.connector
from sqlalchemy import create_engine
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

import urllib.parse

# Configuration
DB_HOST = "localhost"
DB_USER = "root"
DB_PASS = "Root@12345"
ENCODED_PASS = urllib.parse.quote_plus(DB_PASS)
DB_NAME = "ecommerce_analytics"
PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(PROJECT_DIR, "data")
EXCEL_PATH = os.path.join(PROJECT_DIR, "excel", "Ecommerce_Sales_Analytics_Report.xlsx")

def setup_mysql_database():
    print("[1/5] Connecting to MySQL server...")
    conn = mysql.connector.connect(host=DB_HOST, user=DB_USER, password=DB_PASS)
    cursor = conn.cursor()
    
    schema_file = os.path.join(PROJECT_DIR, "sql", "schema.sql")
    with open(schema_file, 'r') as f:
        sql_statements = f.read().split(';')
        
    for statement in sql_statements:
        statement = statement.strip()
        if statement:
            cursor.execute(statement)
            
    conn.commit()
    cursor.close()
    conn.close()
    print("Database 'ecommerce_analytics' and tables created successfully!")

def load_data_to_mysql():
    print("[2/5] Loading CSV data into MySQL database tables...")
    engine = create_engine(f"mysql+mysqlconnector://{DB_USER}:{ENCODED_PASS}@{DB_HOST}/{DB_NAME}")
    
    df_cust = pd.read_csv(os.path.join(DATA_DIR, "raw_customers.csv"))
    df_prod = pd.read_csv(os.path.join(DATA_DIR, "raw_products.csv"))
    df_ord = pd.read_csv(os.path.join(DATA_DIR, "raw_orders.csv"))
    
    df_cust.to_sql("customers", con=engine, if_exists="append", index=False)
    df_prod.to_sql("products", con=engine, if_exists="append", index=False)
    df_ord.to_sql("orders", con=engine, if_exists="append", index=False)
    
    print(f"Data loaded into MySQL: {len(df_cust)} Customers, {len(df_prod)} Products, {len(df_ord)} Orders.")

def run_analytical_queries():
    print("[3/5] Querying analytical metrics from MySQL...")
    engine = create_engine(f"mysql+mysqlconnector://{DB_USER}:{ENCODED_PASS}@{DB_HOST}/{DB_NAME}")
    
    kpi_sql = """
    SELECT 
        COUNT(DISTINCT order_id) AS Total_Orders,
        COUNT(DISTINCT customer_id) AS Unique_Customers,
        SUM(quantity) AS Total_Units_Sold,
        ROUND(SUM(total_revenue), 2) AS Total_Revenue,
        ROUND(SUM(profit), 2) AS Total_Profit,
        ROUND(AVG(total_revenue), 2) AS Average_Order_Value,
        ROUND((SUM(profit) / SUM(total_revenue)) * 100, 2) AS Profit_Margin_Pct
    FROM orders WHERE order_status = 'Delivered';
    """
    
    monthly_sql = """
    SELECT 
        DATE_FORMAT(order_date, '%Y-%m') AS Month,
        COUNT(order_id) AS Total_Orders,
        ROUND(SUM(total_revenue), 2) AS Revenue,
        ROUND(SUM(profit), 2) AS Profit,
        ROUND((SUM(profit) / SUM(total_revenue)) * 100, 2) AS Profit_Margin_Pct
    FROM orders WHERE order_status = 'Delivered'
    GROUP BY DATE_FORMAT(order_date, '%Y-%m') ORDER BY Month ASC;
    """
    
    prod_sql = """
    SELECT 
        p.product_name AS Product,
        p.category AS Category,
        SUM(o.quantity) AS Units_Sold,
        ROUND(SUM(o.total_revenue), 2) AS Total_Revenue,
        ROUND(SUM(o.profit), 2) AS Total_Profit
    FROM orders o JOIN products p ON o.product_id = p.product_id
    WHERE o.order_status = 'Delivered'
    GROUP BY p.product_name, p.category ORDER BY Total_Revenue DESC;
    """
    
    region_sql = """
    SELECT 
        c.state AS State,
        c.city AS City,
        COUNT(o.order_id) AS Total_Orders,
        ROUND(SUM(o.total_revenue), 2) AS Total_Revenue,
        ROUND(SUM(o.profit), 2) AS Total_Profit
    FROM orders o JOIN customers c ON o.customer_id = c.customer_id
    WHERE o.order_status = 'Delivered'
    GROUP BY c.state, c.city ORDER BY Total_Revenue DESC;
    """
    
    df_kpi = pd.read_sql(kpi_sql, con=engine)
    df_monthly = pd.read_sql(monthly_sql, con=engine)
    df_prod = pd.read_sql(prod_sql, con=engine)
    df_region = pd.read_sql(region_sql, con=engine)
    df_orders = pd.read_sql("SELECT * FROM orders;", con=engine)
    df_customers = pd.read_sql("SELECT * FROM customers;", con=engine)
    df_products = pd.read_sql("SELECT * FROM products;", con=engine)
    
    return {
        "Executive Summary": df_kpi,
        "Monthly Trend": df_monthly,
        "Product Performance": df_prod,
        "Regional Breakdown": df_region,
        "orders": df_orders,
        "customers": df_customers,
        "products": df_products
    }

def export_formatted_excel(data_dict):
    print("[4/5] Exporting formatted multi-tab Excel report...")
    writer = pd.ExcelWriter(EXCEL_PATH, engine="openpyxl")
    
    for sheet_name, df in data_dict.items():
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        
    writer.close()
    
    # Apply Professional Styling with openpyxl
    wb = openpyxl.load_workbook(EXCEL_PATH)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for sheet in wb.worksheets:
        sheet.views.sheetView[0].showGridLines = True
        # Style headers
        for col in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Style data cells & auto-adjust column width
        for col in range(1, sheet.max_column + 1):
            max_len = max(len(str(sheet.cell(row=r, column=col).value or '')) for r in range(1, sheet.max_row + 1))
            col_letter = get_column_letter(col)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row=row, column=col)
                cell.font = cell_font
                cell.border = thin_border
                
    wb.save(EXCEL_PATH)
    print(f"[5/5] Excel report saved to: {EXCEL_PATH}")

if __name__ == "__main__":
    setup_mysql_database()
    load_data_to_mysql()
    reports = run_analytical_queries()
    export_formatted_excel(reports)
    print("\nETL & Analytics Pipeline executed successfully!")
